import datetime
import io
import pandas as pd
import streamlit as st
import hashlib

from openpyxl.styles import PatternFill, Border, Side, Alignment, Font
from openpyxl.utils import get_column_letter

import psycopg2
from psycopg2 import IntegrityError
from sqlalchemy import create_engine

# ---------------------------------------------------------
# SETUP & CONFIG
# ---------------------------------------------------------
st.set_page_config(page_title="ระบบวันลา เวชระเบียน", page_icon="🏥", layout="wide")

if "logged_in" not in st.session_state:
    st.session_state.logged_in = False
    st.session_state.emp_id = ""
    st.session_state.emp_name = ""
    st.session_state.role = ""

def make_hash(password):
    return hashlib.sha256(password.encode()).hexdigest()

def to_thai_date(dt_obj):
    months = ["", "มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน", "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม"]
    year = dt_obj.year + 543
    return f"{dt_obj.day} {months[dt_obj.month]} {year}"

def to_thai_date_short(date_str):
    if not date_str: return ""
    months = ["", "ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.", "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.", "พ.ย.", "ธ.ค."]
    dt = datetime.datetime.strptime(date_str, '%Y-%m-%d')
    year = dt.year + 543
    return f"{dt.day} {months[dt.month]} {year}"

def format_date_range(start_str, end_str):
    if not start_str: return ""
    if start_str == end_str:
        return to_thai_date_short(start_str)
        
    dt_start = datetime.datetime.strptime(start_str, '%Y-%m-%d')
    dt_end = datetime.datetime.strptime(end_str, '%Y-%m-%d')
    
    if dt_start.month != dt_end.month or dt_start.year != dt_end.year:
        return f"{to_thai_date_short(start_str)} - {to_thai_date_short(end_str)}"
    
    months = ["", "ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.", "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.", "พ.ย.", "ธ.ค."]
    year = dt_start.year + 543
    return f"{dt_start.day}-{dt_end.day} {months[dt_start.month]} {year}"

def get_month_year_thai(date_str):
    months = ["", "มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน", "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม"]
    dt = datetime.datetime.strptime(date_str, '%Y-%m-%d')
    return f"เดือน {months[dt.month]} {dt.year + 543}"

# ---------------------------------------------------------
# DATABASE SETUP (SUPABASE POSTGRESQL)
# ---------------------------------------------------------
DB_URL = "postgresql://postgres.koslwumkjvvlhboutpcy:GNKCxEPuD11@aws-0-ap-southeast-1.pooler.supabase.com:6543/postgres"
engine = create_engine(DB_URL.replace("postgresql://", "postgresql+psycopg2://"))

def get_connection():
    return psycopg2.connect(DB_URL)

def add_user(emp_id, name, phone, email, password, role='user', q_personal=3.0, q_sick=15.0, q_vacation=10.0):
    conn = get_connection()
    c = conn.cursor()
    try:
        c.execute('''INSERT INTO users (emp_id, name, phone, email, password, role, quota_personal, quota_sick, quota_vacation) 
                     VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)''', 
                  (emp_id.strip(), name.strip(), phone.strip(), email.strip(), make_hash(password), role, q_personal, q_sick, q_vacation))
        conn.commit()
        success = True
    except IntegrityError:
        success = False 
    conn.close()
    return success

def update_user_quota(emp_id, q_personal, q_sick, q_vacation):
    conn = get_connection()
    c = conn.cursor()
    c.execute("UPDATE users SET quota_personal=%s, quota_sick=%s, quota_vacation=%s WHERE emp_id=%s", 
              (q_personal, q_sick, q_vacation, emp_id))
    conn.commit()
    conn.close()

# ฟังก์ชันอัปเดตรหัสผ่าน
def update_user_password(emp_id, new_password):
    conn = get_connection()
    c = conn.cursor()
    c.execute("UPDATE users SET password=%s WHERE emp_id=%s", (make_hash(new_password), emp_id))
    conn.commit()
    conn.close()

def delete_user(emp_id):
    conn = get_connection()
    c = conn.cursor()
    c.execute("DELETE FROM users WHERE emp_id=%s", (emp_id,))
    conn.commit()
    conn.close()

def get_all_users():
    conn = get_connection()
    c = conn.cursor()
    c.execute("SELECT emp_id, name, quota_personal, quota_sick, quota_vacation FROM users WHERE role='user' ORDER BY id ASC")
    users = c.fetchall()
    conn.close()
    return users

def login_user(emp_id, password):
    conn = get_connection()
    c = conn.cursor()
    c.execute("SELECT role, name FROM users WHERE emp_id=%s AND password=%s", (emp_id.strip(), make_hash(password)))
    result = c.fetchone()
    conn.close()
    return result

# คำนวณวันลา (ไม่นับรวม เสาร์-อาทิตย์)
def calculate_leave_days(start_date, start_type, end_date, end_type):
    if start_date > end_date: return 0.0
    total_days = 0.0
    
    if start_date == end_date:
        if start_date.weekday() < 5: # วันจันทร์(0) - วันศุกร์(4)
            if start_type in ["ครึ่งวันแรก (เช้า)", "ครึ่งวันหลัง (บ่าย)"]: return 0.5
            return 1.0
        return 0.0
        
    delta = (end_date - start_date).days + 1
    for i in range(delta):
        curr = start_date + datetime.timedelta(days=i)
        if curr.weekday() < 5: # ถ้าเป็นจันทร์-ศุกร์
            if curr == start_date:
                total_days += 0.5 if start_type in ["ครึ่งวันแรก (เช้า)", "ครึ่งวันหลัง (บ่าย)"] else 1.0
            elif curr == end_date:
                total_days += 0.5 if end_type in ["ครึ่งวันแรก (เช้า)", "ครึ่งวันหลัง (บ่าย)"] else 1.0
            else:
                total_days += 1.0
                
    return total_days

def add_leave_request(emp_id, user_name, leave_type, start_date, start_day_type, end_date, end_day_type, total_days, reason, status='ปกติ'):
    conn = get_connection()
    c = conn.cursor()
    c.execute('''
        INSERT INTO leave_requests (emp_id, user_name, position, approver, leave_type, start_date, start_day_type, end_date, end_day_type, total_days, reason, status)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
    ''', (emp_id, user_name, "-", "-", leave_type, str(start_date), start_day_type, str(end_date), end_day_type, total_days, reason, status))
    conn.commit()
    conn.close()

def cancel_leave_request(req_id):
    conn = get_connection()
    c = conn.cursor()
    c.execute("UPDATE leave_requests SET status='ยกเลิก' WHERE id=%s", (req_id,))
    conn.commit()
    conn.close()

def get_all_leaves(emp_id=None):
    if emp_id:
        df = pd.read_sql_query(f"SELECT * FROM leave_requests WHERE emp_id='{emp_id}' ORDER BY start_date DESC", engine)
    else:
        df = pd.read_sql_query("SELECT * FROM leave_requests ORDER BY start_date DESC", engine)
    return df

def get_leave_balances(emp_id=None):
    users = get_all_users()
    df_leaves = get_all_leaves()
    current_year = datetime.datetime.now().year
    
    if not df_leaves.empty:
        df_leaves['year'] = pd.to_datetime(df_leaves['start_date']).dt.year
        df_leaves_current_year = df_leaves[(df_leaves['year'] == current_year) & (df_leaves['status'] == 'ปกติ')]
    else:
        df_leaves_current_year = pd.DataFrame()
    
    balances = []
    for u in users:
        u_id, u_name, q_per, q_sick, q_vac = u
        if emp_id and u_id != emp_id:
            continue
            
        if not df_leaves_current_year.empty:
            user_leaves = df_leaves_current_year[df_leaves_current_year['emp_id'] == u_id]
            used_personal = user_leaves[user_leaves['leave_type'] == 'ลากิจ']['total_days'].sum()
            used_sick = user_leaves[user_leaves['leave_type'] == 'ลาป่วย']['total_days'].sum()
            used_vacation = user_leaves[user_leaves['leave_type'] == 'ลาพักผ่อน']['total_days'].sum()
        else:
            used_personal = used_sick = used_vacation = 0.0
            
        balances.append({
            "รหัสพนักงาน": u_id,
            "ชื่อ-นามสกุล": u_name,
            "ลากิจ (ใช้/โควต้า)": f"{used_personal} / {q_per}",
            "ลากิจ (เหลือ)": q_per - used_personal,
            "ลาป่วย (ใช้/โควต้า)": f"{used_sick} / {q_sick}",
            "ลาป่วย (เหลือ)": q_sick - used_sick,
            "ลาพักผ่อน (ใช้/โควต้า)": f"{used_vacation} / {q_vac}",
            "ลาพักผ่อน (เหลือ)": q_vac - used_vacation,
            "q_per": q_per, "q_sick": q_sick, "q_vac": q_vac
        })
    return pd.DataFrame(balances)

def delete_old_leave_records(before_year):
    """ลบประวัติการลาที่เกิดขึ้นก่อนปีที่ระบุ (เช่น ระบุ 2024 จะลบของปี 2023 ลงไปทั้งหมด)"""
    conn = get_connection()
    c = conn.cursor()
    # สร้างวันที่ 1 มกราคม ของปีนั้น เพื่อลบข้อมูลที่เก่ากว่า
    target_date = f"{before_year}-01-01"
    c.execute("DELETE FROM leave_requests WHERE start_date < %s", (target_date,))
    deleted_count = c.rowcount
    conn.commit()
    conn.close()
    return deleted_count

# ---------------------------------------------------------
# FRONTEND
# ---------------------------------------------------------
if not st.session_state.logged_in:
    st.title("🔐 เข้าสู่ระบบ (หน่วยเวชระเบียน)")
    with st.container():
        log_empid = st.text_input("รหัสพนักงาน (หรือ 'admin')")
        log_password = st.text_input("รหัสผ่าน", type="password")
        if st.button("เข้าสู่ระบบ", type="primary"):
            user_data = login_user(log_empid, log_password)
            if user_data:
                st.session_state.logged_in = True
                st.session_state.emp_id = log_empid
                st.session_state.emp_name = user_data[1]
                st.session_state.role = user_data[0]
                st.rerun()
            else:
                st.error("รหัสพนักงานหรือรหัสผ่านไม่ถูกต้อง!")
else:
    st.sidebar.title(f"👤 สวัสดี, {st.session_state.emp_name}")
    st.sidebar.caption(f"สถานะ: {'ผู้ดูแลระบบ (Admin)' if st.session_state.role == 'admin' else 'พนักงาน'}")
    if st.sidebar.button("🚪 ออกจากระบบ"):
        st.session_state.logged_in = False
        st.session_state.emp_id = ""
        st.session_state.emp_name = ""
        st.session_state.role = ""
        st.rerun()

    # =========================================================
    # สิทธิ์ ADMIN
    # =========================================================
    if st.session_state.role == "admin":
        st.title("🛠️ หน้าต่างผู้ดูแลระบบ (Admin)")
        tab1, tab2, tab3, tab4, tab5 = st.tabs(["👥 จัดการพนักงาน", "📋 ประวัติ & ยกเลิก", "⏰ ขาด/สาย", "📤 Excel", "🗑️ ล้างประวัติเก่า"])

        with tab1:
            st.subheader("➕ เพิ่มบัญชีพนักงานใหม่")
            col1, col2, col3 = st.columns(3)
            with col1:
                new_empid = st.text_input("รหัสพนักงาน (8 หลัก) *", max_chars=8)
                new_name = st.text_input("ชื่อ-นามสกุล *")
                new_phone = st.text_input("เบอร์โทรศัพท์")
            with col2:
                new_email = st.text_input("อีเมล")
                new_pass = st.text_input("รหัสผ่าน *", type="password")
                new_conf = st.text_input("ยืนยันรหัสผ่าน *", type="password")
            with col3:
                q_v = st.number_input("โควต้า: ลาพักผ่อน", value=10.0, step=0.5)
                q_s = st.number_input("โควต้า: ลาป่วย", value=15.0, step=0.5)
                q_p = st.number_input("โควต้า: ลากิจ", value=3.0, step=0.5)
            
            if st.button("บันทึกพนักงานใหม่", type="primary"):
                if not (new_empid and new_name and new_pass and new_conf):
                    st.warning("⚠️ กรุณากรอกข้อมูลที่มีเครื่องหมาย * ให้ครบ")
                elif new_pass != new_conf:
                    st.error("⚠️ รหัสผ่านไม่ตรงกัน!")
                else:
                    if add_user(new_empid, new_name, new_phone, new_email, new_pass, 'user', q_p, q_s, q_v):
                        st.success(f"✅ เพิ่มพนักงานสำเร็จ!")
                    else:
                        st.error("❌ รหัสพนักงานนี้มีอยู่ในระบบแล้ว!")
            
            st.divider()
            
            col_admin1, col_admin2 = st.columns(2)
            with col_admin1:
                st.subheader("⚙️ ปรับโควต้าวันลาพนักงาน")
                users = get_all_users()
                if users:
                    edit_user = st.selectbox("เลือกพนักงาน", [f"{u[0]} - {u[1]}" for u in users])
                    selected_emp_id = edit_user.split(" - ")[0]
                    current_data = next((u for u in users if u[0] == selected_emp_id), None)
                    
                    ce1, ce2, ce3 = st.columns(3)
                    with ce1: eq_v = st.number_input("พักผ่อน", value=float(current_data[4]), step=0.5)
                    with ce2: eq_s = st.number_input("ป่วย", value=float(current_data[3]), step=0.5)
                    with ce3: eq_p = st.number_input("กิจ", value=float(current_data[2]), step=0.5)
                    
                    if st.button("บันทึกโควต้าใหม่"):
                        update_user_quota(selected_emp_id, eq_p, eq_s, eq_v)
                        st.success("✅ อัปเดตโควต้าเรียบร้อย")
                        st.rerun()

            with col_admin2:
                st.subheader("🔑 รีเซ็ตรหัสผ่าน (กรณีพนักงานลืม)")
                if users:
                    reset_user = st.selectbox("เลือกพนักงานที่ต้องการรีเซ็ตรหัส", [f"{u[0]} - {u[1]}" for u in users], key='reset_u')
                    reset_emp_id = reset_user.split(" - ")[0]
                    new_reset_pass = st.text_input("ตั้งรหัสผ่านชั่วคราวใหม่", type="password")
                    
                    if st.button("รีเซ็ตรหัสผ่าน", type="secondary"):
                        if len(new_reset_pass) < 4:
                            st.error("รหัสผ่านควรมีอย่างน้อย 4 ตัว")
                        else:
                            update_user_password(reset_emp_id, new_reset_pass)
                            st.success(f"เปลี่ยนรหัสผ่านให้ {reset_emp_id} สำเร็จ! แจ้งพนักงานให้ล็อกอินได้เลย")

        with tab2:
            st.subheader("📋 ประวัติการลา และ จัดการสถานะ")
            all_leaves = get_all_leaves()
            if not all_leaves.empty:
                display_df = all_leaves[['id', 'emp_id', 'user_name', 'leave_type', 'start_date', 'end_date', 'total_days', 'status', 'reason']]
                display_df.columns = ["ID", "รหัส", "ชื่อ", "ประเภท", "วันเริ่ม", "วันสิ้นสุด", "จำนวน", "สถานะ", "เหตุผล"]
                st.dataframe(display_df, use_container_width=True, hide_index=True)
                
                st.markdown("**ต้องการยกเลิกวันลา?**")
                cancel_id = st.number_input("ระบุ ID ของรายการที่ต้องการยกเลิก (ดูจากตารางด้านบน)", min_value=0, step=1)
                if st.button("ยกเลิกรายการนี้", type="secondary"):
                    cancel_leave_request(cancel_id)
                    st.success(f"ยกเลิกรายการ ID {cancel_id} สำเร็จ! โควต้าจะถูกคืนให้พนักงาน")
                    st.rerun()
            else:
                st.info("ยังไม่มีข้อมูลการลา")

        with tab3:
            st.subheader("⏰ ลงบันทึก สาย / ไม่ทาบบัตร / ขาดงาน")
            if users:
                col_sp1, col_sp2 = st.columns(2)
                with col_sp1:
                    sp_user = st.selectbox("เลือกพนักงาน", [f"{u[0]} - {u[1]}" for u in users], key='sp_u')
                    sp_date = st.date_input("วันที่เกิดเหตุ")
                with col_sp2:
                    sp_type = st.selectbox("ประเภทเหตุการณ์", ["มาสาย", "ไม่ทาบบัตรเข้า", "ไม่ลงเวลาออก", "ขาดงาน"])
                    sp_remark = st.text_input("หมายเหตุ (เช่น ลืมบัตร, รถติด)")
                
                if st.button("บันทึกเหตุการณ์", type="primary"):
                    e_id, e_name = sp_user.split(" - ")
                    add_leave_request(e_id, e_name, sp_type, sp_date, "เต็มวัน", sp_date, "เต็มวัน", 1.0, sp_remark)
                    st.success(f"บันทึกข้อมูล {sp_type} ของ {e_name} สำเร็จ!")

        with tab4:
            st.subheader("📤 ดาวน์โหลดรายงาน Excel (รูปแบบสมบูรณ์)")
            col_ex1, col_ex2 = st.columns(2)
            with col_ex1:
                export_start = st.date_input("ตั้งแต่วันที่:", datetime.date.today().replace(day=1))
            with col_ex2:
                export_end = st.date_input("ถึงวันที่:", datetime.date.today())
                
            all_data = get_all_leaves()
            
            if not all_data.empty:
                all_data['start_dt'] = pd.to_datetime(all_data['start_date']).dt.date
                all_data['end_dt'] = pd.to_datetime(all_data['end_date']).dt.date
                mask_exp = (all_data['start_dt'] <= export_end) & (all_data['end_dt'] >= export_start)
                df_export = all_data.loc[mask_exp].copy()
                
                if not df_export.empty:
                    st.write(f"พบข้อมูลทั้งหมด **{len(df_export)}** รายการ เตรียมพร้อมสร้าง Excel")
                    
                    buffer = io.BytesIO()
                    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
                        
                        start_thai = to_thai_date(export_start)
                        end_thai = to_thai_date(export_end)
                        
                        yellow_fill = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")
                        white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
                        gray_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid") 
                        green_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
                        orange_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
                        blue_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid") 
                        
                        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
                        font_header_big = Font(size=14, bold=True)
                        font_header_sub = Font(size=12, bold=True)
                        font_table_head = Font(size=11, bold=True)
                        font_table_data = Font(size=11, bold=False)

                        # สรุปภาพรวม
                        summary_data = []
                        users_list = get_all_users()
                        for idx, u in enumerate(users_list, 1):
                            u_id, u_name = u[0], u[1]
                            u_df = df_export[(df_export['emp_id'] == u_id) & (df_export['status'] == 'ปกติ')]
                            
                            vacation = u_df[u_df['leave_type'] == 'ลาพักผ่อน']['total_days'].sum()
                            sick = u_df[u_df['leave_type'] == 'ลาป่วย']['total_days'].sum()
                            personal = u_df[u_df['leave_type'] == 'ลากิจ']['total_days'].sum()
                            absent = u_df[u_df['leave_type'] == 'ขาดงาน']['total_days'].sum()
                            late = u_df[u_df['leave_type'] == 'มาสาย']['total_days'].sum()
                            miss_in = u_df[u_df['leave_type'] == 'ไม่ทาบบัตรเข้า']['total_days'].sum()
                            miss_out = u_df[u_df['leave_type'] == 'ไม่ลงเวลาออก']['total_days'].sum()
                            
                            summary_data.append([
                                idx, u_id, u_name, 
                                vacation if vacation > 0 else "", 
                                sick if sick > 0 else "", 
                                personal if personal > 0 else "", 
                                absent if absent > 0 else "", 
                                late if late > 0 else "", 
                                miss_in if miss_in > 0 else "", 
                                miss_out if miss_out > 0 else "", 
                                "" 
                            ])
                            
                        df_summary = pd.DataFrame(summary_data, columns=["ลำดับ", "ID", "ชื่อ-สกุล", "พักผ่อน", "ป่วย", "กิจ", "ขาดงาน", "สาย", "ไม่ลงเวลาเข้า", "ไม่ลงเวลาออก", "หมายเหตุ"])
                        df_summary.to_excel(writer, index=False, sheet_name="ข้อมูลสรุปวันลา", startrow=3)
                        ws_sum = writer.sheets["ข้อมูลสรุปวันลา"]
                        
                        ws_sum.merge_cells('A1:K1')
                        ws_sum['A1'] = "ข้อมูลสรุปวันลา"
                        ws_sum['A1'].font = font_header_big
                        ws_sum['A1'].alignment = Alignment(horizontal="center", vertical="center")
                        
                        ws_sum.merge_cells('A2:K2')
                        ws_sum['A2'] = f"ตั้งแต่วันที่ {start_thai} - {end_thai}"
                        ws_sum['A2'].font = font_header_sub
                        ws_sum['A2'].alignment = Alignment(horizontal="center", vertical="center")
                        
                        for cell in ws_sum[4]:
                            cell.fill = yellow_fill
                            cell.font = font_table_head
                            cell.border = thin_border
                            cell.alignment = Alignment(horizontal="center", vertical="center")
                            
                        for i, row in enumerate(ws_sum.iter_rows(min_row=5, max_row=ws_sum.max_row, min_col=1, max_col=11)):
                            bg_color = gray_fill if i % 2 == 1 else white_fill
                            for cell in row:
                                cell.fill = bg_color
                                cell.font = font_table_data
                                cell.border = thin_border
                                if cell.column == 3: cell.alignment = Alignment(horizontal="left", vertical="center")
                                else: cell.alignment = Alignment(horizontal="center", vertical="center")
                                    
                        for col_idx in range(1, ws_sum.max_column + 1):
                            col_letter = get_column_letter(col_idx)
                            max_length = 0
                            for row_idx in range(4, ws_sum.max_row + 1): 
                                val = ws_sum.cell(row=row_idx, column=col_idx).value
                                if val is not None:
                                    try:
                                        if len(str(val)) > max_length: max_length = len(str(val))
                                    except: pass
                            ws_sum.column_dimensions[col_letter].width = max_length + 8
                            
                        # แยกตามรายชื่อพนักงาน
                        for idx, u in enumerate(users_list, 1):
                            u_id, u_name, q_per, q_sick, q_vac = u
                            sheet_name = f"{idx} {u_name.split(' ')[0]}"[:31] 
                            
                            u_records = df_export[df_export['emp_id'] == u_id].sort_values(by='start_date')
                            
                            s_vac = s_sick = s_per = s_abs = s_late = s_in = s_out = 0.0
                            person_data = []
                            current_month = ""
                            
                            person_data.append(["โควต้าวันลาสะสม", q_vac, q_sick, q_per, "", "", "", "", ""])
                            
                            for _, r in u_records.iterrows():
                                is_cancel = (r['status'] == 'ยกเลิก')
                                remark = f"[ยกเลิก] {r['reason']}" if is_cancel else r['reason']
                                val = 0 if is_cancel else r['total_days']
                                lt = r['leave_type']
                                
                                r_month = get_month_year_thai(r['start_date'])
                                if r_month != current_month:
                                    person_data.append([r_month, "", "", "", "", "", "", "", ""])
                                    current_month = r_month
                                
                                if not is_cancel:
                                    if lt == 'ลาพักผ่อน': s_vac += val
                                    elif lt == 'ลาป่วย': s_sick += val
                                    elif lt == 'ลากิจ': s_per += val
                                    elif lt == 'ขาดงาน': s_abs += val
                                    elif lt == 'มาสาย': s_late += val
                                    elif lt == 'ไม่ทาบบัตรเข้า': s_in += val
                                    elif lt == 'ไม่ลงเวลาออก': s_out += val

                                person_data.append([
                                    format_date_range(r['start_date'], r['end_date']),
                                    val if lt == 'ลาพักผ่อน' else "",
                                    val if lt == 'ลาป่วย' else "",
                                    val if lt == 'ลากิจ' else "",
                                    val if lt == 'ขาดงาน' else "",
                                    val if lt == 'มาสาย' else "",
                                    val if lt == 'ไม่ทาบบัตรเข้า' else "",
                                    val if lt == 'ไม่ลงเวลาออก' else "",
                                    remark
                                ])
                            
                            person_data.append(["รวมใช้ไป", s_vac if s_vac else "", s_sick if s_sick else "", s_per if s_per else "", s_abs if s_abs else "", s_late if s_late else "", s_in if s_in else "", s_out if s_out else "", ""])
                            person_data.append(["คงเหลือ", q_vac - s_vac, q_sick - s_sick, q_per - s_per, "", "", "", "", ""])
                            
                            df_person = pd.DataFrame(person_data, columns=["วันที่", "พักผ่อน", "ป่วย", "กิจ", "ขาดงาน", "สาย", "ไม่ลงเวลาเข้า", "ไม่ลงเวลาออก", "หมายเหตุ"])
                            df_person.to_excel(writer, index=False, sheet_name=sheet_name, startrow=3)
                            
                            ws_p = writer.sheets[sheet_name]
                            
                            ws_p.merge_cells('A1:I1')
                            ws_p['A1'] = f"ประวัติการลา: {u_name} (รหัสพนักงาน: {u_id})"
                            ws_p['A1'].font = font_header_big
                            ws_p['A1'].alignment = Alignment(horizontal="center", vertical="center")
                            
                            ws_p.merge_cells('A2:I2')
                            ws_p['A2'] = f"ตั้งแต่วันที่ {start_thai} - {end_thai}"
                            ws_p['A2'].font = font_header_sub
                            ws_p['A2'].alignment = Alignment(horizontal="center", vertical="center")
                            
                            for cell in ws_p[4]:
                                cell.fill = yellow_fill
                                cell.font = font_table_head
                                cell.border = thin_border
                                cell.alignment = Alignment(horizontal="center", vertical="center")
                                
                            max_r = ws_p.max_row
                            month_row_indices = []

                            for i, row in enumerate(ws_p.iter_rows(min_row=5, max_row=max_r, min_col=1, max_col=9)):
                                is_quota_row = (row[0].row == 5)
                                is_summary_row = (row[0].row >= max_r - 1)
                                val_str = str(row[0].value) if row[0].value else ""
                                is_month_row = val_str.startswith("เดือน ") and (row[1].value == "" or row[1].value is None)
                                
                                if is_summary_row: bg_color = orange_fill; f_style = font_table_head
                                elif is_quota_row: bg_color = green_fill; f_style = font_table_head
                                elif is_month_row: bg_color = blue_fill; f_style = font_table_head; month_row_indices.append(row[0].row)
                                else: bg_color = gray_fill if i % 2 == 1 else white_fill; f_style = font_table_data
                                    
                                for cell in row:
                                    cell.fill = bg_color
                                    cell.font = f_style
                                    cell.border = thin_border
                                    if cell.column == 9 and not is_month_row: cell.alignment = Alignment(horizontal="left", vertical="center")
                                    elif cell.column == 1 and (is_summary_row or is_quota_row): cell.alignment = Alignment(horizontal="right", vertical="center")
                                    elif is_month_row and cell.column == 1: cell.alignment = Alignment(horizontal="left", vertical="center")
                                    else: cell.alignment = Alignment(horizontal="center", vertical="center")
                            
                            for r_idx in month_row_indices:
                                ws_p.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx, end_column=9)
                            
                            for col_idx in range(1, ws_p.max_column + 1):
                                col_letter = get_column_letter(col_idx)
                                max_length = 0
                                for row_idx in range(4, ws_p.max_row + 1):
                                    val = ws_p.cell(row=row_idx, column=col_idx).value
                                    if val is not None:
                                        try:
                                            if len(str(val)) > max_length: max_length = len(str(val))
                                        except: pass
                                ws_p.column_dimensions[col_letter].width = max_length + 8
                    
                    st.download_button(
                        label="📥 ดาวน์โหลดไฟล์ Excel (สรุป + แยกรายบุคคล)",
                        data=buffer.getvalue(),
                        file_name=f"Leave_Report_{export_start}_to_{export_end}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        type="primary"
                    )

        with tab5:
            st.subheader("🗑️ ล้างข้อมูลประวัติการลาเก่า (ลดพื้นที่ฐานข้อมูล)")
            st.warning("⚠️ **คำเตือน:** การลบข้อมูลจะไม่สามารถกู้คืนได้ แนะนำให้ใช้เมนู Export Excel เก็บประวัติรายปีไว้ในคอมพิวเตอร์ก่อนกดลบ")
            
            current_year = datetime.datetime.now().year
            
            # ให้เลือกปีที่ต้องการเคลียร์ทิ้ง
            clear_year = st.number_input("ลบประวัติการลาที่เกิดขึ้น 'ก่อน' ปี (ค.ศ.):", 
                                         min_value=2020, 
                                         max_value=current_year, 
                                         value=current_year, 
                                         step=1)
            
            st.info(f"📌 ระบบจะลบประวัติการลาทั้งหมดที่เกิดขึ้นตั้งแต่ **31 ธันวาคม {clear_year-1} ลงไป** (พนักงานทุกคนยังอยู่ปกติ)")
            
            # ระบบป้องกันคนเผลอกดโดน
            confirm_text = st.text_input("พิมพ์คำว่า 'ยืนยันการลบ' ในช่องด้านล่างเพื่อปลดล็อคปุ่ม")
            
            if st.button("🗑️ ยืนยันลบข้อมูลประวัติเก่า", type="primary", disabled=(confirm_text != "ยืนยันการลบ")):
                count = delete_old_leave_records(clear_year)
                if count > 0:
                    st.success(f"✅ ลบข้อมูลประวัติเก่าจำนวน {count} รายการ สำเร็จ!")
                    st.balloons()
                else:
                    st.info("ไม่พบข้อมูลประวัติที่เก่ากว่าปีที่ระบุ (ไม่มีการลบเกิดขึ้น)")

    # =========================================================
    # สิทธิ์ USER (คนลา)
    # =========================================================
    elif st.session_state.role == "user":
        st.title("📝 ระบบวันลา (หน่วยเวชระเบียน)")
        tab1, tab2, tab3 = st.tabs(["📤 ยื่นใบลา", "📋 ประวัติและสิทธิวันลา", "⚙️ ตั้งค่าบัญชี"]) # เพิ่ม Tab ตั้งค่า
        
        with tab1:
            df_leaves = get_all_leaves()
            with st.container():
                st.markdown("### 1. ประเภทการลา")
                leave_type = st.radio("เลือกประเภทการลาที่ต้องการ", ["ลากิจ", "ลาป่วย", "ลาพักผ่อน"], horizontal=True)
                st.divider()

                st.markdown("### 2. ข้อมูลผู้ขอลา")
                st.text_input("ชื่อ-นามสกุล", value=f"{st.session_state.emp_id} - {st.session_state.emp_name}", disabled=True)
                st.divider()

                st.markdown("### 3. ระยะเวลาที่ต้องการลา (ระบบจะไม่คำนวณวันเสาร์-อาทิตย์ให้)")
                col3_1, col3_2 = st.columns(2)
                with col3_1:
                    start_date = st.date_input("ขอลาตั้งแต่วันที่", datetime.date.today())
                    start_day_type = st.selectbox("รูปแบบวันเริ่มลา", ["เต็มวัน", "ครึ่งวันแรก (เช้า)", "ครึ่งวันหลัง (บ่าย)"])
                with col3_2:
                    end_date = st.date_input("ถึงวันที่", datetime.date.today())
                    end_day_type = st.selectbox("รูปแบบวันสิ้นสุด", ["เต็มวัน", "ครึ่งวันแรก (เช้า)", "ครึ่งวันหลัง (บ่าย)"])

                st.divider()
                reason = st.text_area("4. วัตถุประสงค์ของการลา (โปรดระบุให้ชัดเจน)")
                
                can_leave = True
                if start_date > end_date:
                    st.error("⚠️ วันเริ่มต้นต้องไม่เกินวันสิ้นสุด")
                    can_leave = False
                else:
                    req_days = calculate_leave_days(start_date, start_day_type, end_date, end_day_type)
                    if req_days <= 0:
                        st.error("⚠️ ช่วงวันที่ท่านเลือก ตรงกับวันหยุดเสาร์-อาทิตย์ หรือคำนวณวันลาได้ 0 วัน")
                        can_leave = False
                    else:
                        my_balances = get_leave_balances(st.session_state.emp_id)
                        if not my_balances.empty:
                            rem_days = my_balances.iloc[0][f"{leave_type} (เหลือ)"]
                            if req_days > rem_days:
                                st.error(f"❌ สิทธิการ**{leave_type}**ของท่านในปีนี้ คงเหลือเพียง {rem_days} วัน (ต้องการใช้ {req_days} วัน)")
                                can_leave = False

                        if can_leave and not df_leaves.empty:
                            delta = end_date - start_date
                            requested_dates = [start_date + datetime.timedelta(days=i) for i in range(delta.days + 1)]
                            blocked_dates = []
                            warning_dates = []

                            for d in requested_dates:
                                if d.weekday() >= 5: continue # ไม่เช็คเสาร์อาทิตย์ว่าใครลาชนไหม
                                count_on_day = 0
                                for _, row in df_leaves.iterrows():
                                    if row['status'] == 'ยกเลิก': continue
                                    e_start = datetime.datetime.strptime(row['start_date'], "%Y-%m-%d").date()
                                    e_end = datetime.datetime.strptime(row['end_date'], "%Y-%m-%d").date()
                                    if e_start <= d <= e_end and row['leave_type'] in ['ลากิจ', 'ลาป่วย', 'ลาพักผ่อน']:
                                        count_on_day += 1
                                if count_on_day >= 2:
                                    can_leave = False
                                    blocked_dates.append(f"- วันที่ {d.strftime('%d/%m/%Y')} (โควต้าเต็ม 2 คนแล้ว)")
                                elif count_on_day == 1:
                                    warning_dates.append(f"- วันที่ {d.strftime('%d/%m/%Y')} (มีเจ้าหน้าที่ลาแล้ว 1 คน)")
                            
                            if not can_leave:
                                st.error("❌ **ไม่สามารถยื่นใบลาได้ เนื่องจากมีผู้ลาเต็มโควต้าในวันที่ดังต่อไปนี้:**\n" + "\n".join(blocked_dates))
                            else:
                                if warning_dates:
                                    st.warning("⚠️ **ข้อควรทราบ:** ในช่วงวันที่ท่านเลือก มีผู้ลาไปแล้ว 1 คน (หากท่านกดยืนยัน โควตาวันดังกล่าวจะเต็ม 2 คนพอดี)\n" + "\n".join(warning_dates))
                                st.success(f"✅ วันที่ท่านเลือก สามารถยื่นลาได้ (คำนวณวันลา: {req_days} วัน)")

                if st.button("📤 ส่งใบลา", type="primary", use_container_width=True, disabled=not can_leave):
                    if not reason.strip():
                        st.warning("⚠️ กรุณาระบุวัตถุประสงค์ของการลา")
                    else:
                        add_leave_request(st.session_state.emp_id, st.session_state.emp_name, leave_type, start_date, start_day_type, end_date, end_day_type, req_days, reason)
                        st.success("🎉 บันทึกข้อมูลการลาเข้าสู่ระบบเรียบร้อยแล้ว!")
                        st.balloons()
                        
        with tab2:
            st.subheader(f"📊 สิทธิวันลาคงเหลือของฉัน (ประจำปี {datetime.datetime.now().year})")
            my_balances = get_leave_balances(st.session_state.emp_id)
            if not my_balances.empty:
                col_b1, col_b2, col_b3 = st.columns(3)
                data = my_balances.iloc[0]
                with col_b1: st.metric("ลากิจ (เหลือ)", f"{data['ลากิจ (เหลือ)']} วัน", f"จาก {data['q_per']} วัน", delta_color="off")
                with col_b2: st.metric("ลาป่วย (เหลือ)", f"{data['ลาป่วย (เหลือ)']} วัน", f"จาก {data['q_sick']} วัน", delta_color="off")
                with col_b3: st.metric("ลาพักผ่อน (เหลือ)", f"{data['ลาพักผ่อน (เหลือ)']} วัน", f"จาก {data['q_vac']} วัน", delta_color="off")
            
            st.divider()
            st.subheader("📋 ประวัติการลา และ ข้อมูลเข้างานของฉัน")
            my_leaves = get_all_leaves(st.session_state.emp_id)
            if not my_leaves.empty:
                display_my = my_leaves[['id', 'leave_type', 'start_date', 'end_date', 'total_days', 'status', 'reason']]
                display_my.columns = ["ID", "ประเภท", "วันเริ่ม", "วันสิ้นสุด", "จำนวน", "สถานะ", "เหตุผล"]
                st.dataframe(display_my, use_container_width=True, hide_index=True)
                
                # ฟีเจอร์ให้พนักงานยกเลิกใบลาของตัวเองได้
                st.markdown("**ต้องการยกเลิกวันลาของตนเอง?**")
                cancel_my_id = st.number_input("ระบุ ID ของใบลาที่ต้องการยกเลิก", min_value=0, step=1)
                if st.button("ยกเลิกใบลา", type="secondary"):
                    df_check = my_leaves[(my_leaves['id'] == cancel_my_id) & (my_leaves['status'] == 'ปกติ')]
                    if not df_check.empty:
                        # ดึงวันที่เริ่มลามาเช็ค
                        start_date_str = df_check.iloc[0]['start_date']
                        leave_start_date = datetime.datetime.strptime(start_date_str, "%Y-%m-%d").date()
                        
                        # เช็คว่าวันเริ่มลา น้อยกว่าหรือเท่ากับ วันนี้หรือไม่ (ถ้าเป็นอดีตหรือวันนี้ จะไม่ให้ยกเลิกเอง)
                        if leave_start_date <= datetime.date.today():
                            st.error("❌ ไม่สามารถยกเลิกวันลาที่ผ่านมาแล้ว (หรือวันลาที่กำลังใช้งานอยู่ในวันนี้) ได้ หากมีความจำเป็นกรุณาติดต่อ Admin ครับ")
                        else:
                            cancel_leave_request(cancel_my_id)
                            st.success("✅ ยกเลิกใบลาสำเร็จ โควตาวันลาของคุณถูกคืนให้แล้ว!")
                            st.rerun()
                    else:
                        st.error("❌ ไม่พบข้อมูล ID นี้ หรือรายการนี้ถูกยกเลิกไปแล้ว")
            else:
                st.info("ยังไม่มีประวัติในระบบ")
                
        # Tab 3: สำหรับเปลี่ยนรหัสผ่าน
        with tab3:
            st.subheader("🔑 เปลี่ยนรหัสผ่าน")
            new_pw = st.text_input("รหัสผ่านใหม่", type="password")
            confirm_pw = st.text_input("ยืนยันรหัสผ่านใหม่", type="password")
            
            if st.button("บันทึกรหัสผ่านใหม่", type="primary"):
                if not new_pw or not confirm_pw:
                    st.warning("กรุณากรอกข้อมูลให้ครบถ้วน")
                elif new_pw != confirm_pw:
                    st.error("รหัสผ่านไม่ตรงกัน!")
                elif len(new_pw) < 4:
                    st.error("รหัสผ่านควรมีอย่างน้อย 4 ตัวอักษร")
                else:
                    update_user_password(st.session_state.emp_id, new_pw)
                    st.success("✅ อัปเดตรหัสผ่านเรียบร้อยแล้ว!")
